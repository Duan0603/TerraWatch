"""
Script sinh bộ Slide Thuyết trình PowerPoint (10 trang) chuẩn Đồ án Tốt nghiệp
Chuyên ngành: Kỹ thuật Phần mềm (Software Engineering Capstone)
Đề tài: GEOSENTRY (TERRAWATCH) — HỆ THỐNG VIỄN THÁM & AI CẢNH BÁO SỚM SẠT LỞ ĐẤT
"""

import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# -----------------------------------------------------------------------------
# BẢNG MÀU CHUYÊN NGHIỆP DARK THEME (HIGH-TECH DISASTER RISK MONITORING)
# -----------------------------------------------------------------------------
BG_COLOR = RGBColor(11, 19, 32)       # Xanh đen đậm sang trọng (#0B1320)
CARD_BG = RGBColor(19, 30, 49)        # Nền card (#131E31)
CARD_BORDER = RGBColor(31, 58, 82)    # Viền card (#1F3A52)
TEXT_WHITE = RGBColor(255, 255, 255)  # Trắng thuần
TEXT_MUTED = RGBColor(156, 178, 201)  # Xám xanh nhạt
CYAN_ACCENT = RGBColor(0, 212, 255)   # Xanh ngọc công nghệ (#00D4FF)
ORANGE_ACCENT = RGBColor(255, 140, 0) # Cam cảnh báo (#FF8C00)
GREEN_ACCENT = RGBColor(16, 185, 129) # Xanh lá thành công (#10B981)
RED_ACCENT = RGBColor(239, 68, 68)    # Đỏ nguy cấp (#EF4444)

def add_slide_background(slide):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_COLOR
    bg.line.fill.background()
    return bg

def add_header(slide, title_text, category="GEOSENTRY • CAPSTONE DEFENSE 2026"):
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.1))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p_cat = tf.paragraphs[0]
    p_cat.text = category.upper()
    p_cat.font.size = Pt(10.5)
    p_cat.font.bold = True
    p_cat.font.color.rgb = ORANGE_ACCENT

    p_title = tf.add_paragraph()
    p_title.text = title_text
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_WHITE
    p_title.space_before = Pt(3)

def add_card(slide, left, top, width, height, title, bullet_points, accent_color=CYAN_ACCENT, subtitle=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD_BG
    shape.line.color.rgb = CARD_BORDER
    shape.line.width = Pt(1)

    bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent_color
    bar.line.fill.background()

    tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.18), width - Inches(0.4), height - Inches(0.25))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p_title = tf.paragraphs[0]
    p_title.text = title
    p_title.font.size = Pt(13)
    p_title.font.bold = True
    p_title.font.color.rgb = accent_color

    if subtitle:
        p_sub = tf.add_paragraph()
        p_sub.text = subtitle
        p_sub.font.size = Pt(9.5)
        p_sub.font.color.rgb = TEXT_MUTED
        p_sub.space_after = Pt(4)

    for pt in bullet_points:
        p = tf.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(3)

def create_geosentry_deck(output_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # -------------------------------------------------------------------------
    # SLIDE 1: TITLE SLIDE
    # -------------------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    add_slide_background(s1)

    tb = s1.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.7), Inches(2.8))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "ĐỒ ÁN TỐT NGHIỆP KỸ SƯ PHẦN MỀM (CAPSTONE PROJECT)"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = ORANGE_ACCENT

    p1 = tf.add_paragraph()
    p1.text = "GEOSENTRY (TERRAWATCH)"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p1.space_before = Pt(6)

    p2 = tf.add_paragraph()
    p2.text = "Hệ Thống Viễn Thám & Trí Tuệ Nhân Tạo Cảnh Báo Sớm Sạt Lở Đất Miền Núi"
    p2.font.size = Pt(18)
    p2.font.color.rgb = CYAN_ACCENT
    p2.space_before = Pt(8)

    # Team Members Box
    t_box = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.2), Inches(11.7), Inches(2.2))
    t_box.fill.solid()
    t_box.fill.fore_color.rgb = CARD_BG
    t_box.line.color.rgb = CARD_BORDER

    tf_team = t_box.text_frame
    tf_team.margin_top = Inches(0.2)
    tf_team.margin_left = Inches(0.4)
    p_h = tf_team.paragraphs[0]
    p_h.text = "NHÓM SINH VIÊN THỰC HIỆN ĐỀ TÀI (5 THÀNH VIÊN WBS):"
    p_h.font.size = Pt(12)
    p_h.font.bold = True
    p_h.font.color.rgb = ORANGE_ACCENT

    members = [
        "1. Duẫn — AI / Computer Vision Engineer (Model DeepLabV3+ ONNX, Tensor 8 kênh, Xếp hạng mức nguy cơ)",
        "2. Tú — GIS Pipeline & Data Engineer (Sentinel-2 Crawl, Cloud Masking, NDVI, DEM Slope, PostGIS)",
        "3. Thuận — Backend Engineer & System Architect (Spring Boot 3, JWT 3 Roles, Phát cảnh báo SMS/Push, Gateway)",
        "4. Huy — Frontend WebGIS Engineer (React 18 + Vite, Mapbox 3D Terrain, Hàng đợi duyệt, Nút Cảnh báo Admin)",
        "5. Lâm — Mobile App Engineer (Flutter 3.x, SQLite Offline 10,000 đa giác, Geofencing Hú Còi, Nhận cảnh báo Push/SMS)"
    ]
    for m in members:
        p_m = tf_team.add_paragraph()
        p_m.text = m
        p_m.font.size = Pt(10.5)
        p_m.font.color.rgb = TEXT_WHITE
        p_m.space_before = Pt(2)

    # -------------------------------------------------------------------------
    # SLIDE 2: PROBLEM STATEMENT
    # -------------------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_slide_background(s2)
    add_header(s2, "TÍNH CẤP THIẾT & 4 ĐIỂM NGHẼN TRONG CẢNH BÁO SẠT LỞ")

    card_w = Inches(2.7)
    card_h = Inches(4.8)
    top_pos = Inches(1.8)

    add_card(s2, Inches(0.8), top_pos, card_w, card_h, "1. Phát Hiện Chậm Trễ", [
        "Sạt lở xảy ra bất ngờ sau mưa dầm kéo dài tại vùng núi hiểm trở.",
        "Điển hình: Vụ sạt lở Làng Nủ (Lào Cai) sau bão Yagi 2024.",
        "Hiện tại chỉ phát hiện khi thảm họa đã ập xuống hoặc có người dân chạy bộ về báo.",
        "Thiếu cơ chế quan trắc diện rộng tự động liên tục."
    ], RED_ACCENT, "Phụ thuộc báo tin thủ công")

    add_card(s2, Inches(3.8), top_pos, card_w, card_h, "2. Mất Sóng Khi Mưa Bão", [
        "Mưa bão quật đổ cột BTS, sạt đường đứt cáp viễn thông.",
        "Các bản làng bị cô lập mất sạch 100% sóng 4G/Internet.",
        "Người dân không thể nhận tin tức qua mạng khi di chuyển qua vùng nguy hiểm.",
        "Ứng dụng thông thường hoàn toàn tê liệt khi không có Internet."
    ], ORANGE_ACCENT, "Mất kết nối hoàn toàn")

    add_card(s2, Inches(6.8), top_pos, card_w, card_h, "3. Bản Đồ 2D Thiếu Trực Quan", [
        "Bản đồ 2D phẳng không thể hiện được độ dốc sườn núi hiểm trở.",
        "Không quan sát được vết trượt bùn đất đang lan rộng theo hướng nào.",
        "Cán bộ khó đánh giá tương quan giữa vách dốc và khu dân cư.",
        "Khó đối chiếu ảnh vệ tinh trước và sau để thẩm định nguy cơ."
    ], CYAN_ACCENT, "Hạn chế của bản đồ phẳng")

    add_card(s2, Inches(9.8), top_pos, card_w, card_h, "4. Cảnh Báo Thiếu Kịp Thời", [
        "Thiếu công cụ cho người có thẩm quyền chủ động phát cảnh báo khẩn cấp.",
        "Chưa tích hợp đồng thời kênh SMS trực tiếp và Push Notification di động.",
        "Chưa có cơ chế lọc người nhận chính xác theo khu vực chịu ảnh hưởng.",
        "Người dân không nhận được chỉ thị sơ tán khẩn cấp trước giờ G."
    ], GREEN_ACCENT, "Thiếu kênh cảnh báo tức thì")

    # -------------------------------------------------------------------------
    # SLIDE 3: THE SOLUTION
    # -------------------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_slide_background(s3)
    add_header(s3, "GIẢI PHÁP GEOSENTRY: HỆ SINH THÁI CẢNH BÁO SỚM SẠT LỞ ĐỒNG BỘ")

    card_w2 = Inches(5.6)
    card_h2 = Inches(2.3)

    add_card(s3, Inches(0.8), Inches(1.8), card_w2, card_h2, "🛰️ Tầng Viễn Thám Vĩ Mô (Sentinel-2 + DEM)", [
        "Vệ tinh quét diện rộng không cần thiết bị hiện trường.",
        "Lọc mây thông minh và tính chỉ số suy giảm thảm phủ Delta-NDVI.",
        "Trích xuất độ dốc sườn đồi từ ảnh số độ cao SRTM 30m DEM."
    ], CYAN_ACCENT)

    add_card(s3, Inches(6.9), Inches(1.8), card_w2, card_h2, "🧠 Tầng Phân Tích AI & Xếp Hạng Nguy Cơ", [
        "DeepLabV3+ ONNX nhận diện đa giác vết trượt sạt lở (F1 = 0.768).",
        "Tự động vector hóa ra chuẩn PostGIS MultiPolygon GeoJSON.",
        "Tự động xếp hạng nguy cơ theo độ dốc và khoảng cách khu dân cư."
    ], GREEN_ACCENT)

    add_card(s3, Inches(0.8), Inches(4.5), card_w2, card_h2, "🖥️ Tầng Chỉ Huy & Cảnh Báo (WebGIS Command Center)", [
        "Bản đồ địa hình Mapbox 3D Terrain sườn núi sống động.",
        "Thanh trượt đối soát ảnh trước/sau và hàng đợi duyệt sạt lở 1-click.",
        "Nút '🚨 Phát Cảnh Báo Khẩn Cấp' cho Admin gửi SMS + App Push."
    ], ORANGE_ACCENT)

    add_card(s3, Inches(6.9), Inches(4.5), card_w2, card_h2, "📱 Tầng Ngoại Tuyến Hiện Trường (Mobile Citizen App)", [
        "Cơ chế Geofencing Ngoại tuyến: SQLite lưu 10,000 đa giác sạt lở.",
        "Rung chuông còi hú báo động âm lượng tối đa ngay khi mất sóng 4G.",
        "Đăng ký nhận cảnh báo khẩn cấp và gửi báo cáo hiện trường."
    ], RED_ACCENT)

    # -------------------------------------------------------------------------
    # SLIDE 4: ARCHITECTURE (8 COMPONENTS)
    # -------------------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_slide_background(s4)
    add_header(s4, "KIẾN TRÚC MICROSERVICES 8 THÀNH PHẦN CỐT LÕI")

    grid_w = Inches(2.7)
    grid_h = Inches(2.3)

    add_card(s4, Inches(0.8), Inches(1.8), grid_w, grid_h, "1. Clients", [
        "WebGIS React 18 + Mapbox GL",
        "Mobile App Flutter 3.x + SQLite",
        "Thiết kế Responsive & Dark Theme"
    ], CYAN_ACCENT)

    add_card(s4, Inches(3.8), Inches(1.8), grid_w, grid_h, "2. API Gateway", [
        "Nginx Reverse Proxy (Port 8080)",
        "Rate Limiting 50 req/s",
        "Global CORS & Gzip Compression"
    ], CYAN_ACCENT)

    add_card(s4, Inches(6.8), Inches(1.8), grid_w, grid_h, "3. Core API Service", [
        "Java 17 + Spring Boot 3",
        "Spring Security 6 + JWT 3 Roles",
        "Duyệt sạt lở & Phát cảnh báo khẩn cấp"
    ], GREEN_ACCENT)

    add_card(s4, Inches(9.8), Inches(1.8), grid_w, grid_h, "4. AI & GIS Services", [
        "AI: FastAPI + ONNX Runtime",
        "GIS: FastAPI + GDAL/Rasterio",
        "Tách biệt Bounded Context"
    ], GREEN_ACCENT)

    add_card(s4, Inches(0.8), Inches(4.5), grid_w, grid_h, "5. PostGIS Database", [
        "PostgreSQL 15 + PostGIS 3.3",
        "Logical Schema-per-Service",
        "core_schema & gis_schema GiST"
    ], ORANGE_ACCENT)

    add_card(s4, Inches(3.8), Inches(4.5), grid_w, grid_h, "6. Message Broker", [
        "Redis 7 Alpine Pub/Sub",
        "Channel: terrawatch:events",
        "Xử lý bất đồng bộ đa service"
    ], ORANGE_ACCENT)

    add_card(s4, Inches(6.8), Inches(4.5), grid_w, grid_h, "7. Fault Tolerance", [
        "Resilience4j Circuit Breaker",
        "Ngắt mạch khi AI/GIS lỗi >= 50%",
        "Fallback queue, Core API không sập"
    ], RED_ACCENT)

    add_card(s4, Inches(9.8), Inches(4.5), grid_w, grid_h, "8. Kênh Phát Cảnh Báo", [
        "Firebase Cloud Messaging (FCM Push)",
        "SMS Gateway (eSMS / SpeedSMS / Twilio)",
        "Gửi đồng thời về App và số điện thoại"
    ], RED_ACCENT)

    # -------------------------------------------------------------------------
    # SLIDE 5: TU (GIS PIPELINE)
    # -------------------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_slide_background(s5)
    add_header(s5, "PHÂN HỆ GIS PIPELINE & VIỄN THÁM (TÚ PHỤ TRÁCH)", "THÀNH VIÊN 2 • GIS PIPELINE & DATA ENGINEER")

    c_w = Inches(3.7)
    c_h = Inches(4.8)

    add_card(s5, Inches(0.8), Inches(1.8), c_w, c_h, "Thu Thập & Lọc Mây Vệ Tinh", [
        "Tự động crawl ảnh Sentinel-2 L2A qua Copernicus CDSE / Sentinel Hub API.",
        "Áp dụng thuật toán lọc mây thông minh dựa trên Scene Classification Layer (SCL / s2cloudless).",
        "Tự động loại bỏ các khung cảnh bị mây che phủ > 40%.",
        "Lưu trữ siêu dữ liệu cảnh chụp vào gis_schema.satellite_scenes."
    ], CYAN_ACCENT, "Copernicus API & Cloud Masking")

    add_card(s5, Inches(4.8), Inches(1.8), c_w, c_h, "Tính Toán Chỉ Số Địa Không Gian", [
        "Tính toán biến động thực vật Delta-NDVI giữa cảnh trước và sau thiên tai bằng rasterio.",
        "Trích xuất góc dốc địa hình sườn núi [0 - 90 độ] từ ảnh số độ cao SRTM 30m DEM.",
        "Chuẩn hóa và đồng bộ hệ quy chiếu tọa độ địa lý WGS84 EPSG:4326.",
        "Chuẩn bị ma trận 8 kênh hoàn chỉnh bàn giao cho phân hệ AI."
    ], GREEN_ACCENT, "Delta-NDVI & SRTM 30m DEM Slope")

    add_card(s5, Inches(8.8), Inches(1.8), c_w, c_h, "Tiling & Raster-to-Vector", [
        "Module Tiling băm lưới khu vực giám sát AOI thành các patch 128x128 pixel.",
        "Băm lưới thông minh tính toán độ nén kinh độ theo vĩ độ và độ gối mép (stride overlap).",
        "Thuật toán Raster-to-Vector (Polygonization) chuyển mặt nạ AI thành đa giác PostGIS MultiPolygon.",
        "Phục vụ Vector Tiles MVT qua /tiles/{z}/{x}/{y}.pbf cho WebGIS."
    ], ORANGE_ACCENT, "Tiling Engine & Vector Tiles")

    # -------------------------------------------------------------------------
    # SLIDE 6: DUAN (AI / CV)
    # -------------------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    add_slide_background(s6)
    add_header(s6, "PHÂN HỆ TRÍ TUỆ NHÂN TẠO & HỌC SÂU (DUẪN PHỤ TRÁCH)", "THÀNH VIÊN 1 • AI / COMPUTER VISION ENGINEER")

    # Metrics row
    m_box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(0.9))
    m_box.fill.solid()
    m_box.fill.fore_color.rgb = CARD_BG
    m_box.line.color.rgb = CYAN_ACCENT
    tf_m = m_box.text_frame
    p = tf_m.paragraphs[0]
    p.text = "CHỈ SỐ ĐÁNH GIÁ MÔ HÌNH:   F1-Score: 0.768   |   IoU: 0.642   |   Độ trễ: < 300ms   |   Đầu vào: Tensor 8 dải phổ"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT
    p.alignment = PP_ALIGN.CENTER

    add_card(s6, Inches(0.8), Inches(2.9), c_w, Inches(3.9), "Semantic Segmentation Sạt Lở", [
        "Kiến trúc DeepLabV3+ với Backbone ResNet-50 / EfficientNet kết hợp ASPP.",
        "Huấn luyện trên tập dữ liệu chuẩn quốc tế Landslide4Sense (> 3,799 mẫu).",
        "Đầu vào 8 kênh: [B2, B3, B4, B8, B11, B12, NDVI, SLOPE].",
        "Hàm mất mát Focal + Dice Loss giải quyết mất cân bằng mẫu nghiêm trọng (< 2% diện tích sạt lở)."
    ], CYAN_ACCENT, "DeepLabV3+ on Landslide4Sense")

    add_card(s6, Inches(4.8), Inches(2.9), c_w, Inches(3.9), "Đóng Gói ONNX Runtime", [
        "Xuất mô hình PyTorch sang định dạng ONNX Runtime tối ưu.",
        "Tốc độ suy luận CPU Intel chỉ 182ms, GPU RTX 3060 đạt 26ms/patch.",
        "Triển khai qua FastAPI service độc lập tại port 8001.",
        "Hỗ trợ chế độ chạy nền tự động xử lý hàng loạt khi có ảnh vệ tinh mới."
    ], GREEN_ACCENT, "High Performance Inference")

    add_card(s6, Inches(8.8), Inches(2.9), c_w, Inches(3.9), "Xếp Hạng Mức Nguy Cơ Sạt Lở", [
        "Tự động tính toán điểm nguy cơ dựa trên: độ dốc trung bình, diện tích khối trượt và độ tin cậy AI.",
        "Truy vấn không gian PostGIS ST_Distance tính khoảng cách tới khu dân cư gần nhất.",
        "Tự động gán nhãn risk_level: low, medium, high, extreme.",
        "Ưu tiên đẩy các điểm sạt lở có nguy cơ đe dọa khu dân cư lên đầu hàng đợi thẩm định."
    ], ORANGE_ACCENT, "Automated Risk Scoring (FR2.3)")

    # -------------------------------------------------------------------------
    # SLIDE 7: THUAN (BACKEND)
    # -------------------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    add_slide_background(s7)
    add_header(s7, "PHÂN HỆ BACKEND, SECURITY & PHÁT CẢNH BÁO (THUẬN PHỤ TRÁCH)", "THÀNH VIÊN 3 • BACKEND ENGINEER & ARCHITECT")

    add_card(s7, Inches(0.8), Inches(1.8), c_w, c_h, "Core API & Bảo Mật RBAC JWT", [
        "Phát triển trên nền tảng Java 17 + Spring Boot 3 công nghiệp.",
        "Bảo mật Spring Security 6 với JWT Token phi trạng thái (Stateless).",
        "Phân quyền RBAC 3 vai trò: citizen, officer, admin.",
        "Băm mật khẩu bằng BCrypt; endpoint /api/v1/auth/** phục vụ đăng ký, đăng nhập an toàn."
    ], CYAN_ACCENT, "Spring Boot 3 & Security 6")

    add_card(s7, Inches(4.8), Inches(1.8), c_w, c_h, "Resilience4j Circuit Breaker", [
        "Bọc các kết nối gọi liên service sang AI và GIS bằng Circuit Breaker.",
        "Cấu hình tự động chuyển sang OPEN khi tỷ lệ lỗi vượt quá 50% trong 10 calls.",
        "Kích hoạt phương thức Fallback lưu request vào hàng đợi xử lý ngầm.",
        "Core API được bảo vệ tuyệt đối, không bao giờ bị sập lan truyền (Cascading Failures)."
    ], RED_ACCENT, "Chống sập lan truyền")

    add_card(s7, Inches(8.8), Inches(1.8), c_w, c_h, "Phát Cảnh Báo Khẩn Cấp & Audit", [
        "API phát cảnh báo khẩn cấp: Gửi đồng thời qua SMS Gateway và Firebase Cloud Messaging.",
        "Lọc người nhận theo vùng quan tâm đăng ký (alert_subscriptions) qua truy vấn PostGIS ST_DWithin.",
        "Lưu vết kiểm toán bất biến (Audit Trail) trong bảng landslide_event_history và alert_broadcasts.",
        "Message Broker Redis 7 Pub/Sub xử lý gửi tin nhắn bất đồng bộ theo lô."
    ], GREEN_ACCENT, "SMS & FCM Alert Broadcast")

    # -------------------------------------------------------------------------
    # SLIDE 8: HUY (WEBGIS 3D)
    # -------------------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    add_slide_background(s8)
    add_header(s8, "PHÂN HỆ WEBGIS COMMAND CENTER & CẢNH BÁO (HUY PHỤ TRÁCH)", "THÀNH VIÊN 4 • FRONTEND WEBGIS ENGINEER")

    add_card(s8, Inches(0.8), Inches(1.8), c_w, c_h, "Mapbox GL 3D Terrain", [
        "Phát triển bằng React 18 + Vite kết hợp Tailwind CSS.",
        "Bật lớp bản đồ địa hình 3D (Terrain 3D Elevation) trên nền WebGL.",
        "Mô phỏng chân thực độ dốc núi rừng hiểm trở các tỉnh Yên Bái, Lào Cai.",
        "Cho phép cán bộ nghiêng, xoay camera 3D 360 độ để quan sát vết nứt sườn đồi."
    ], CYAN_ACCENT, "Địa hình 3D sườn núi chân thực")

    add_card(s8, Inches(4.8), Inches(1.8), c_w, c_h, "Đối Soát Ảnh & Thẩm Định", [
        "Tính năng Time-slider swipe: Kéo thanh trượt đối soát ảnh vệ tinh trước và sau sạt lở.",
        "Cán bộ trực ban dễ dàng xác minh sự biến mất của thảm thực vật bằng mắt thường.",
        "Hàng đợi thẩm định sạt lở 1-click: Xem thông tin diện tích, độ dốc, mức nguy cơ và duyệt nhanh.",
        "Hiển thị các điểm báo cáo hiện trường từ người dân để đối chiếu thực tế."
    ], GREEN_ACCENT, "Time-slider Swipe & 1-Click Approval")

    add_card(s8, Inches(8.8), Inches(1.8), c_w, c_h, "Nút Phát Cảnh Báo Khẩn Cấp (Admin)", [
        "Nút bấm nổi bật '🚨 Phát Cảnh Báo Khẩn Cấp' chỉ hiển thị với vai trò Admin.",
        "Chọn phạm vi vùng nguy hiểm theo điểm sạt lở (vùng đệm bán kính) hoặc toàn vùng giám sát.",
        "Tùy chọn kênh gửi: SMS về số điện thoại, Push Notification về ứng dụng di động, hoặc cả hai.",
        "Xem trước số người nhận và xác nhận 2 bước an toàn chống bấm nhầm."
    ], ORANGE_ACCENT, "Emergency Broadcast Button")

    # -------------------------------------------------------------------------
    # SLIDE 9: LAM (MOBILE GEOFENCING)
    # -------------------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    add_slide_background(s9)
    add_header(s9, "PHÂN HỆ MOBILE CITIZEN APP & GEOFENCING NGOẠI TUYẾN (LÂM PHỤ TRÁCH)", "THÀNH VIÊN 5 • MOBILE APP ENGINEER")

    add_card(s9, Inches(0.8), Inches(1.8), c_w, c_h, "Geofencing Ngoại Tuyến Đột Phá", [
        "Ứng dụng di động Flutter 3.x đa nền tảng (Android & iOS).",
        "Bộ nhớ đệm SQLite cục bộ lưu sẵn >= 10,000 đa giác vùng nguy cơ sạt lở (< 15MB).",
        "Dịch vụ định vị chạy nền (Background Geolocation) liên tục quét tọa độ GPS mỗi 10 giây.",
        "Thuật toán Haversine và Ray-Casting (Point-in-Polygon) kiểm tra không gian siêu tốc (< 20ms)."
    ], CYAN_ACCENT, "Offline-First SQLite & Ray-Casting")

    add_card(s9, Inches(4.8), Inches(1.8), c_w, c_h, "Còi Hú Báo Động Khẩn Cấp", [
        "Tự động kích hoạt khi người dân di chuyển vào vùng sạt lở nguy hiểm.",
        "HOẠT ĐỘNG 100% NGOẠI TUYẾN ngay cả khi mất sạch sóng 4G/Internet.",
        "Chớp màn hình đỏ cảnh báo, rung máy liên tục và hú còi âm lượng tối đa.",
        "Có cơ chế bypass chế độ im lặng của điện thoại để đánh thức người dân trong đêm."
    ], RED_ACCENT, "Báo động khẩn khi mất mạng")

    add_card(s9, Inches(8.8), Inches(1.8), c_w, c_h, "Nhận Cảnh Báo & Báo Cáo Hiện Trường", [
        "Nhận thông báo khẩn cấp toàn màn hình từ Admin qua Firebase Cloud Messaging (FCM).",
        "Đăng ký vùng quan tâm và cập nhật số điện thoại nhận tin nhắn SMS cảnh báo.",
        "Module Báo cáo hiện trường: Chụp ảnh thực địa và gửi tọa độ GPS vết nứt/trượt đất.",
        "Lưu hàng đợi cục bộ tự động đồng bộ khi có lại kết nối mạng."
    ], GREEN_ACCENT, "Push/SMS Alert & Field Reports")

    # -------------------------------------------------------------------------
    # SLIDE 10: ROADMAP & CONCLUSION
    # -------------------------------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    add_slide_background(s10)
    add_header(s10, "LỘ TRÌNH TRIỂN KHAI, ĐỊNH HƯỚNG NÂNG CAO & KẾT LUẬN")

    card_w3 = Inches(5.6)
    card_h3 = Inches(2.3)

    add_card(s10, Inches(0.8), Inches(1.8), card_w3, card_h3, "🚀 Kế Hoạch 5 Tuần Tốc Lực (MVP)", [
        "Tuần 1: Dựng nền tảng, PostGIS & Auth JWT 3 roles.",
        "Tuần 2: Ingestion Sentinel-2, DEM Slope và API báo cáo hiện trường.",
        "Tuần 3: Tích hợp mô hình ONNX thật, cắt ảnh Tiling và xếp hạng nguy cơ.",
        "Tuần 4: Dashboard WebGIS thẩm định, Nút Cảnh Báo Khẩn Cấp Admin và Mobile FCM.",
        "Tuần 5: Thông luồng 100%, kiểm thử tải và demo kịch bản thực tế Làng Nủ."
    ], CYAN_ACCENT)

    add_card(s10, Inches(6.9), Inches(1.8), card_w3, card_h3, "📦 Nâng Cấp Capstone (Hướng phát triển Epic 6)", [
        "Trạm IoT ESP32 quan trắc rung chấn sườn dốc (cảm biến rung & độ ẩm đất đẩy MQTT).",
        "Tích hợp ảnh Radar SAR Sentinel-1 quan sát xuyên mây trong mùa mưa bão.",
        "Tích hợp chỉ số mưa tích lũy Antecedent Rainfall Index (ARI) từ vệ tinh GPM NASA.",
        "Chạy thử nghiệm trên tập dữ liệu thực tế Lào Cai / Yên Bái sau bão Yagi."
    ], ORANGE_ACCENT)

    # Conclusion Banner
    concl = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.5), Inches(11.7), Inches(2.2))
    concl.fill.solid()
    concl.fill.fore_color.rgb = RGBColor(16, 40, 30)
    concl.line.color.rgb = GREEN_ACCENT
    concl.line.width = Pt(1.5)

    tf_c = concl.text_frame
    tf_c.margin_left = Inches(0.4)
    tf_c.margin_top = Inches(0.25)
    p_c1 = tf_c.paragraphs[0]
    p_c1.text = "🎯 KẾT LUẬN & GIÁ TRỊ THỰC TIỄN CỦA ĐỒ ÁN GEOSENTRY:"
    p_c1.font.size = Pt(14)
    p_c1.font.bold = True
    p_c1.font.color.rgb = GREEN_ACCENT

    points = [
        "1. Giải quyết bài toán cấp bách mang tính nhân văn sâu sắc: Giám sát, phát hiện và cảnh báo sớm thiên tai sạt lở đất bảo vệ người dân.",
        "2. Hiện thực hóa kiến trúc Microservices 8 thành phần công nghiệp, kết hợp Viễn thám Sentinel-2, AI ONNX Runtime và WebGL 3D.",
        "3. Tích hợp nút Phát Cảnh Báo Khẩn Cấp cho Admin: Chủ động cảnh báo người dân qua cả SMS và thông báo đẩy ứng dụng di động.",
        "4. Đột phá với cơ chế Geofencing Ngoại tuyến: Đảm bảo người dân vẫn được còi hú cứu mạng ngay cả khi mất sạch 100% sóng 4G/Internet."
    ]
    for pt in points:
        p_pt = tf_c.add_paragraph()
        p_pt.text = pt
        p_pt.font.size = Pt(11)
        p_pt.font.color.rgb = TEXT_WHITE
        p_pt.space_before = Pt(3)

    # Save
    prs.save(output_path)
    print(f"Successfully generated custom presentation: {output_path}")

if __name__ == "__main__":
    out = "d:/SECapstone/docs/GeoSentry_Capstone_Presentation.pptx"
    if len(sys.argv) > 1:
        out = sys.argv[1]
    create_geosentry_deck(out)
